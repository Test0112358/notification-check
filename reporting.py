#!/usr/bin/env python3
"""
Reports, spreadsheets and emails for the ACCC Acquisitions Register Monitor
===========================================================================
Everything a person reads is produced here, from data the scraper has
already saved:

  data/register.csv        flat export (Google Sheets / Excel import)
  data/register.xlsx       formatted workbook (emailed)
  data/summary.json        statistics, deadlines, analytics, data health
  data/latest_changes.json this run's changes in readable form
  data/status.csv          "Last Updated" tab content
  outbox/*.html            HTML email bodies for the workflow's email steps

Day counts follow the ACCC convention: the notification date is Day 0, so
the first business day counted is the day after it.
"""

import csv
import datetime
import html
import os
import re
import statistics

import documents
from accc_common import (
    HOLIDAY_TABLE_LAST_YEAR, LATEST_CHANGES_JSON, OUTBOX_DIR, REGISTER_CSV,
    REGISTER_URL, REGISTER_XLSX, SCHEDULE_TEXT, STATUS_CSV, SUMMARY_JSON,
    WATCHLIST_JSON, canberra_today, clean, count_accc_business_days, data_age_days,
    display_date, fmt_canberra, iso_date, load_json, parse_date, parse_iso,
    save_json, utc_now,
)
from anzsic_sectors import get_ma_sectors_for_entry

# ---------------------------------------------------------------------------
# Labels and formatting helpers
# ---------------------------------------------------------------------------
FIELD_LABELS = {
    "title": "Matter name",
    "acquisition_status": "Status",
    "type": "Type",
    "case_number": "Case number",
    "stage": "Stage",
    "notification_date": "Notification date",
    "end_of_determination_period": "End of determination period",
    "accc_determination": "ACCC determination",
    "determination_publication_date": "Determination published",
    "acquirers": "Acquirer(s)",
    "targets": "Target(s) / vendor(s)",
    "other_parties": "Other parties",
    "anzsic_codes": "ANZSIC codes",
    "last_modified_accc": "ACCC page last modified",
    "consultation_text": "Consultation text",
    "consultation_close_date": "Consultation closes",
    "consultation_docs": "Consultation documents",
    "decisions_docs": "Decisions and key events",
}

DATE_FIELDS = {"notification_date", "end_of_determination_period",
               "determination_publication_date", "consultation_close_date"}

DUE_SOON_BUSINESS_DAYS = 10
CONSULTATION_REMINDER_BUSINESS_DAYS = 5
RECENT_CHANGES_DAYS = 30


def esc(value):
    return html.escape(str(value if value is not None else ""))


def short(value, limit=160):
    value = clean(value)
    return value if len(value) <= limit else value[: limit - 1].rstrip() + "…"


def _parse_doc_list(value):
    """Changelog stores document lists as their Python repr; recover the rows."""
    rows = []
    for m in re.finditer(r"'date': '([^']*)', 'description': '((?:[^'\\]|\\.)*)', 'url': '([^']*)'", value or ""):
        rows.append({"date": m.group(1), "description": m.group(2), "url": m.group(3)})
    return rows


def describe_change(ch):
    """One readable line for a field change, e.g. 'Status: Under assessment → Assessment completed'."""
    field = ch.get("field", "")
    label = FIELD_LABELS.get(field, field.replace("_", " ").capitalize())
    old, new = ch.get("old", ""), ch.get("new", "")
    if field in ("consultation_docs", "decisions_docs"):
        added = documents.new_documents(_parse_doc_list(old), _parse_doc_list(new))
        if added:
            return "; ".join(
                f"New document: {d['description']} ({display_date(d['date'])})" for d in added)
        return f"{label}: document links updated"
    if field in DATE_FIELDS:
        old, new = display_date(old), display_date(new)
    if field == "last_modified_accc":
        old, new = fmt_canberra(old), fmt_canberra(new)
    old = short(old, 120) or "(blank)"
    new = short(new, 120) or "(removed)"
    return f"{label}: {old} → {new}"


def repo_url(path=""):
    server = os.environ.get("GITHUB_SERVER_URL", "https://github.com")
    repo = os.environ.get("GITHUB_REPOSITORY", "")
    if not repo:
        return ""
    return f"{server}/{repo}{path}"


def run_url():
    run_id = os.environ.get("GITHUB_RUN_ID", "")
    return repo_url(f"/actions/runs/{run_id}") if run_id else repo_url("/actions")


# ---------------------------------------------------------------------------
# Deadline tracker
# ---------------------------------------------------------------------------

def deadline_info(entry, today=None):
    """
    Where an active matter sits in its determination period.
    Returns dict(day, of, remaining, status, label) — empty values when not applicable.
    """
    today = today or canberra_today()
    info = {"day": "", "of": "", "remaining": "", "status": "", "label": ""}
    status = entry.get("acquisition_status", "")
    notif = parse_date(entry.get("notification_date", ""))
    end = parse_date(entry.get("end_of_determination_period", ""))

    if status and status != "Under assessment":
        info["status"] = "Decided" if entry.get("accc_determination") else status
        return info
    if not notif:
        return info
    if not end:
        info["status"] = "Awaiting end date" if entry.get("type") == "Notification" else "Under assessment"
        return info

    period = count_accc_business_days(notif, end) or 0
    elapsed = count_accc_business_days(notif, min(today, end)) or 0
    remaining = count_accc_business_days(today, end) if today < end else 0
    info.update({"day": elapsed, "of": period, "remaining": remaining})
    if today > end:
        info["status"] = "Past deadline (awaiting publication)"
    elif today == end:
        info["status"] = "Due today"
    elif remaining <= 5:
        info["status"] = "Due within 5 business days"
    elif remaining <= DUE_SOON_BUSINESS_DAYS:
        info["status"] = "Due within 10 business days"
    else:
        info["status"] = "On track"
    info["label"] = f"Day {elapsed} of {period}"
    return info


def consultation_info(entry, today=None):
    """(close_date or None, business days left or None)."""
    today = today or canberra_today()
    close = parse_date(entry.get("consultation_close_date", ""))
    if not close or close < today:
        return close, None
    return close, count_accc_business_days(today, close)


def latest_key_event(entry):
    docs = entry.get("decisions_docs") or []
    if not docs:
        return ""
    dated = sorted(docs, key=lambda d: parse_date(d.get("date")) or datetime.date.min)
    d = dated[-1]
    category, _ = documents.classify_decision_document(d.get("description", ""))
    return f"{display_date(d.get('date', ''))}: {category}"


# ---------------------------------------------------------------------------
# Watchlist
# ---------------------------------------------------------------------------

def load_watchlist():
    wl = load_json(WATCHLIST_JSON, {})
    return {k: [clean(x) for x in wl.get(k, []) if clean(x)]
            for k in ("companies", "sectors", "anzsic_codes", "keywords")}


def watchlist_matches(entry, watchlist):
    """List of human-readable reasons this matter matches the watchlist."""
    from scraper import party_names  # local import: scraper imports this module
    reasons = []
    names = " | ".join(party_names(entry.get("acquirers", "")) + party_names(entry.get("targets", ""))
                       + party_names(entry.get("other_parties", "")))
    haystack = f"{entry.get('title', '')} | {names}".lower()
    for company in watchlist.get("companies", []):
        if re.search(rf"(?<!\w){re.escape(company.lower())}(?!\w)", haystack):
            reasons.append(f"company: {company}")
    sectors = get_ma_sectors_for_entry(entry.get("anzsic_codes", ""))
    for sector in watchlist.get("sectors", []):
        if any(sector.lower() == s.lower() for s in sectors):
            reasons.append(f"sector: {sector}")
    codes = re.findall(r"\b\d{4}\b", entry.get("anzsic_codes", ""))
    for code in watchlist.get("anzsic_codes", []):
        if any(c.startswith(code) for c in codes):
            reasons.append(f"ANZSIC: {code}")
    text = " ".join([haystack, entry.get("consultation_text", "").lower()]
                    + [d.get("description", "").lower() for d in entry.get("decisions_docs", []) or []])
    for kw in watchlist.get("keywords", []):
        if kw.lower() in text:
            reasons.append(f"keyword: {kw}")
    return reasons


# ---------------------------------------------------------------------------
# Decision document alerts
# ---------------------------------------------------------------------------

def collect_decision_alerts(this_run_events, stored_register, register):
    """New rows in 'Decisions and key events' during this run, classified."""
    alerts = []
    for event in this_run_events:
        if event["event"] not in ("NEW_ENTRY", "CHANGED"):
            continue
        slug = event["slug"]
        entry = register.get(slug, {})
        old_docs = stored_register.get(slug, {}).get("decisions_docs", []) if event["event"] == "CHANGED" else []
        for doc in documents.new_documents(old_docs, entry.get("decisions_docs", [])):
            category, priority = documents.classify_decision_document(doc.get("description", ""))
            alerts.append({
                "slug": slug,
                "title": entry.get("title", slug),
                "case_number": entry.get("case_number", ""),
                "type": entry.get("type", ""),
                "accc_url": entry.get("url", ""),
                "date": doc.get("date", ""),
                "description": doc.get("description", ""),
                "url": doc.get("url", ""),
                "category": category,
                "priority": priority,
                "new_matter": event["event"] == "NEW_ENTRY",
            })
    alerts.sort(key=lambda a: (not a["priority"], a["category"], a["title"]))
    return alerts


# ---------------------------------------------------------------------------
# Analytics
# ---------------------------------------------------------------------------
APPROVED = {"Approved", "Approved with conditions"}
DETERMINED = APPROVED | {"Not approved"}


def _rate(n, d):
    return round(100.0 * n / d, 1) if d else None


def _median(values):
    return round(statistics.median(values), 1) if values else None


def _mean(values):
    return round(statistics.mean(values), 1) if values else None


def _outcome_stats(entries):
    determined = [e for e in entries if e.get("accc_determination") in DETERMINED]
    approved = [e for e in determined if e.get("accc_determination") in APPROVED]
    days = [e["elapsed_business_days"] for e in determined if isinstance(e.get("elapsed_business_days"), int)]
    return {
        "determined": len(determined),
        "approved": sum(1 for e in determined if e.get("accc_determination") == "Approved"),
        "approved_with_conditions": sum(1 for e in determined if e.get("accc_determination") == "Approved with conditions"),
        "not_approved": sum(1 for e in determined if e.get("accc_determination") == "Not approved"),
        "ceased": sum(1 for e in entries if e.get("acquisition_status") == "Assessment ceased"),
        "approval_rate_pct": _rate(len(approved), len(determined)),
        "median_business_days": _median(days),
        "mean_business_days": _mean(days),
        "fastest_business_days": min(days) if days else None,
        "slowest_business_days": max(days) if days else None,
    }


def build_analytics(register):
    entries = list(register.values())
    by_type = {}
    for kind in ("Notification", "Waiver"):
        subset = [e for e in entries if e.get("type") == kind]
        by_type[kind] = {"matters": len(subset), **_outcome_stats(subset)}

    # Month tables across every year present (no hardcoded year)
    months = {}
    for e in entries:
        kind = e.get("type", "Unknown")
        filed = parse_date(e.get("notification_date"))
        decided = parse_date(e.get("determination_publication_date"))
        if filed:
            row = months.setdefault((filed.year, filed.month, kind), {"filed": 0, "decided": []})
            row["filed"] += 1
        if decided:
            row = months.setdefault((decided.year, decided.month, kind), {"filed": 0, "decided": []})
            row["decided"].append(e)
    monthly = []
    for (year, month, kind), row in sorted(months.items()):
        stats = _outcome_stats(row["decided"])
        monthly.append({"year": year, "month": month, "type": kind, "filed": row["filed"],
                        "completed": len(row["decided"]), **stats})

    years = sorted({m["year"] for m in monthly})
    yearly = []
    for year in years:
        for kind in ("Notification", "Waiver"):
            filed = [e for e in entries if e.get("type") == kind
                     and (parse_date(e.get("notification_date")) or datetime.date.min).year == year]
            decided = [e for e in entries if e.get("type") == kind
                       and (parse_date(e.get("determination_publication_date")) or datetime.date.min).year == year]
            yearly.append({"year": year, "type": kind, "filed": len(filed),
                           "completed": len(decided), **_outcome_stats(decided)})

    sectors = {}
    for e in entries:
        for sector in get_ma_sectors_for_entry(e.get("anzsic_codes", "")) or ["Unclassified"]:
            sectors.setdefault(sector, []).append(e)
    by_sector = sorted(
        ({"sector": s, "matters": len(es),
          "under_assessment": sum(1 for e in es if e.get("acquisition_status") == "Under assessment"),
          **_outcome_stats(es)} for s, es in sectors.items()),
        key=lambda r: -r["matters"])

    phases = {}
    for e in entries:
        if e.get("type") == "Notification" and e.get("stage"):
            phases.setdefault(e["stage"], []).append(e)
    by_stage = {stage: {"matters": len(es), **_outcome_stats(es)} for stage, es in phases.items()}

    return {"by_type": by_type, "yearly": yearly, "monthly": monthly,
            "by_sector": by_sector, "by_stage": by_stage}


# ---------------------------------------------------------------------------
# Summary JSON
# ---------------------------------------------------------------------------

def deadline_lists(register, today=None):
    today = today or canberra_today()
    due, overdue, closing = [], [], []
    for slug, e in register.items():
        info = deadline_info(e, today)
        if info["status"] == "Past deadline (awaiting publication)":
            overdue.append((slug, e, info))
        elif info["remaining"] != "" and info["remaining"] <= DUE_SOON_BUSINESS_DAYS:
            due.append((slug, e, info))
        close, left = consultation_info(e, today)
        if left is not None and left <= CONSULTATION_REMINDER_BUSINESS_DAYS:
            closing.append((slug, e, close, left))
    due.sort(key=lambda x: x[2]["remaining"])
    overdue.sort(key=lambda x: parse_date(x[1].get("end_of_determination_period")))
    closing.sort(key=lambda x: x[2])
    return due, overdue, closing


def generate_summary(register, health, manifest):
    """Produce aggregate statistics for the dashboard / README."""
    by_status, by_stage, by_type, by_determination = {}, {}, {}, {}
    for e in register.values():
        s = e.get("acquisition_status") or "Unknown"
        by_status[s] = by_status.get(s, 0) + 1
        st = e.get("stage", "")
        if st:
            by_stage[st] = by_stage.get(st, 0) + 1
        t = e.get("type") or "Unknown"
        by_type[t] = by_type.get(t, 0) + 1
        d = e.get("accc_determination", "")
        if d:
            by_determination[d] = by_determination.get(d, 0) + 1

    due, overdue, closing = deadline_lists(register)
    return {
        "generated_at_utc": utc_now().replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "data_health": {
            "last_success_utc": health.get("last_success_utc", ""),
            "data_age_days": data_age_days(health),
            "consecutive_failures": health.get("consecutive_failures", 0),
            "last_request_count": health.get("last_request_count", 0),
        },
        "total_entries": len(register),
        "by_acquisition_status": by_status,
        "by_stage": by_stage,
        "by_type": by_type,
        "by_accc_determination": by_determination,
        "decisions_due_within_10_days": [
            {"slug": slug, "title": e.get("title", ""), "case_number": e.get("case_number", ""),
             "end_of_determination_period": e.get("end_of_determination_period", ""),
             "business_days_remaining": info["remaining"], "day_of_period": info["label"]}
            for slug, e, info in due
        ],
        "past_deadline_awaiting_publication": [
            {"slug": slug, "title": e.get("title", ""), "case_number": e.get("case_number", ""),
             "end_of_determination_period": e.get("end_of_determination_period", "")}
            for slug, e, info in overdue
        ],
        "consultations_closing_soon": [
            {"slug": slug, "title": e.get("title", ""), "case_number": e.get("case_number", ""),
             "closes": close.isoformat(), "business_days_left": left}
            for slug, e, close, left in closing
        ],
        "documents": documents.manifest_counts(manifest),
        "analytics": build_analytics(register),
    }


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def format_docs_for_csv(docs):
    """
    Format a list of document dicts as a pipe-separated readable string.
    Example: "3 Apr 2026: Statement of Issues | 15 May 2026: Determination"
    """
    if not docs:
        return ""
    return " | ".join(
        f"{display_date(d.get('date', ''))}: {d.get('description', '')}"
        for d in docs
        if d.get("description")
    )


def _int_or_blank(v):
    return v if isinstance(v, int) else ""


def has_questionnaire(entry, manifest_by_slug):
    if any(rec.get("kind") == "Questionnaire" for rec in manifest_by_slug.get(entry.get("slug"), [])):
        return True
    if any("questionnaire" in (d.get("description", "") + d.get("url", "")).lower()
           for d in entry.get("consultation_docs", []) or []):
        return True
    # Consultation tables disappear once consultation closes; the text remains
    return "questionnaire" in entry.get("consultation_text", "").lower()


def register_rows(register, manifest, watchlist, today=None):
    """Rows for the CSV / Sheets raw data. Original columns first, new ones appended."""
    today = today or canberra_today()
    by_slug = {}
    for rec in manifest.values():
        by_slug.setdefault(rec.get("slug"), []).append(rec)
    rows = []
    for slug, e in register.items():
        info = deadline_info(e, today)
        close, _ = consultation_info(e, today)
        rows.append({
            "Acquisition Name":                         e.get("title", ""),
            "Type":                                     e.get("type", ""),
            "Case Number":                              e.get("case_number", ""),
            "Acquirer(s)":                              e.get("acquirers", ""),
            "Target(s) / Vendor(s)":                    e.get("targets", ""),
            "M&A Sector":                               "; ".join(get_ma_sectors_for_entry(e.get("anzsic_codes", ""))),
            "Other Parties":                            e.get("other_parties", ""),
            "Notification / Application Date":          iso_date(e.get("notification_date", "")),
            "Stage":                                    e.get("stage", ""),
            "Acquisition Status":                       e.get("acquisition_status", ""),
            "End of Determination Period":              iso_date(e.get("end_of_determination_period", "")),
            # Calculated span: notification date (Day 0) → end of determination period
            "Determination Period Calendar Days":       _int_or_blank(e.get("det_period_calendar_days")),
            "Determination Period Business Days (ACT)": _int_or_blank(e.get("det_period_business_days")),
            # For completed cases: notification date (Day 0) → determination publication date
            "Elapsed Calendar Days to Decision":        _int_or_blank(e.get("elapsed_calendar_days")),
            "Elapsed Business Days to Decision (ACT)":  _int_or_blank(e.get("elapsed_business_days")),
            "ACCC Determination":                       e.get("accc_determination", ""),
            "Determination Publication Date":           iso_date(e.get("determination_publication_date", "")),
            "ANZSIC Code(s)":                           e.get("anzsic_codes", ""),
            "URL":                                      e.get("url", ""),
            "Decisions and Key Events":                 format_docs_for_csv(e.get("decisions_docs", [])),
            "Consultation Text":                        e.get("consultation_text", ""),
            "Questionnaire":                            "Y" if has_questionnaire(e, by_slug) else "",
            "Consultation Documents":                   format_docs_for_csv(e.get("consultation_docs", [])),
            "ACCC Page Last Modified":                  fmt_canberra(e.get("last_modified_accc", "")),
            "Last Checked (Canberra time)":             fmt_canberra(e.get("last_checked_utc", "")),
            # ── Added columns (appended so existing column letters are unchanged) ──
            "Business Days Elapsed":                    info["day"],
            "Business Days Remaining":                  info["remaining"],
            "Deadline Status":                          info["status"],
            "Consultation Closes":                      close.isoformat() if close else "",
            "Latest Key Event":                         latest_key_event(e),
            "Watchlist Match":                          "; ".join(watchlist_matches(e, watchlist)),
            "First Seen":                               fmt_canberra(e.get("first_seen_utc", ""), "%d %b %Y"),
        })
    rows.sort(key=lambda r: r["Notification / Application Date"], reverse=True)
    return rows


def export_csv(rows):
    """Write the register data to a CSV with human-readable column names."""
    os.makedirs(os.path.dirname(REGISTER_CSV) or ".", exist_ok=True)
    fieldnames = list(rows[0].keys()) if rows else []
    # utf-8-sig BOM so Excel opens it correctly without a manual encoding step
    with open(REGISTER_CSV, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    print(f"  CSV exported: {len(rows)} rows → {REGISTER_CSV}")


# ---------------------------------------------------------------------------
# Case history
# ---------------------------------------------------------------------------

def case_history_rows(changelog, register, since_days=None):
    """Material changelog events as readable rows, newest first."""
    cutoff = utc_now() - datetime.timedelta(days=since_days) if since_days else None
    rows = []
    for event in changelog:
        if not event.get("material", True):
            continue
        when = parse_iso(event.get("timestamp_utc"))
        if cutoff and (not when or when < cutoff):
            continue
        entry = register.get(event["slug"], {})
        case = event.get("case_number") or entry.get("case_number", "")
        base = {
            "when": when,
            "Detected (Canberra time)": fmt_canberra(when),
            "Case Number": case,
            "Matter": event.get("title", event["slug"]),
            "URL": event.get("url") or entry.get("url", ""),
        }
        if event["event"] == "NEW_ENTRY":
            rows.append({**base, "Event": "New matter", "Detail": f"{entry.get('type', '')} added to the register"})
        elif event["event"] == "REMOVED":
            rows.append({**base, "Event": "Removed", "Detail": "No longer on the register"})
        else:
            for ch in event.get("changes", []):
                if ch.get("material", True):
                    rows.append({**base, "Event": FIELD_LABELS.get(ch["field"], ch["field"]),
                                 "Detail": describe_change(ch)})
    rows.sort(key=lambda r: (r["when"] or datetime.datetime.min.replace(tzinfo=datetime.timezone.utc)),
              reverse=True)
    return rows


# ---------------------------------------------------------------------------
# Excel workbook
# ---------------------------------------------------------------------------
NAVY = "0D1B2A"
BLUE = "1D4ED8"
STATUS_FILLS = {
    "Past deadline (awaiting publication)": ("FEE2E2", "991B1B"),
    "Due today":                            ("FEE2E2", "991B1B"),
    "Due within 5 business days":           ("FEF3C7", "92400E"),
    "Due within 10 business days":          ("FEF9C3", "854D0E"),
    "On track":                             ("ECFDF5", "065F46"),
}
DETERMINATION_FONTS = {"Not approved": "991B1B", "Approved with conditions": "92400E"}


def export_excel(register, rows, changelog, analytics, health, run_utc, counts):
    """
    Formatted workbook: Summary, Active Matters, Notifications, Waivers,
    Recent Changes, Case History, Analytics, Raw Data.
    Dates are real Excel dates so they sort and filter correctly.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    header_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    body_font = Font(name="Calibri", size=10)
    link_font = Font(name="Calibri", size=10, color=BLUE, underline="single")
    alt_fill = PatternFill("solid", fgColor="F3F6FA")
    border = Border(bottom=Side(style="thin", color="E5E7EB"))
    wrap = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top")
    date_fmt = "d mmm yyyy"

    def write_table(ws, columns, data, tab_color, widths=None, wrap_cols=(), date_cols=(),
                    link_cols=(), status_col=None, det_col=None):
        ws.sheet_properties.tabColor = tab_color
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 30
        for ci, col in enumerate(columns, 1):
            c = ws.cell(row=1, column=ci, value=col)
            c.fill, c.font, c.border = header_fill, header_font, border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(ci)].width = (widths or {}).get(col, 16)
        for ri, record in enumerate(data, 2):
            fill = alt_fill if ri % 2 == 0 else None
            for ci, col in enumerate(columns, 1):
                val = record.get(col, "")
                if col in date_cols:
                    val = parse_date(val) or (val or None)
                c = ws.cell(row=ri, column=ci, value=val if val != "" else None)
                c.font, c.border = body_font, border
                c.alignment = wrap if col in wrap_cols else top
                if fill:
                    c.fill = fill
                if col in date_cols and isinstance(val, datetime.date):
                    c.number_format = date_fmt
                if col in link_cols and val:
                    c.value, c.hyperlink, c.font = "Open ↗", val, link_font
                if col == status_col and val in STATUS_FILLS:
                    bg, fg = STATUS_FILLS[val]
                    c.fill = PatternFill("solid", fgColor=bg)
                    c.font = Font(name="Calibri", size=10, bold=True, color=fg)
                if col == det_col and val in DETERMINATION_FONTS:
                    c.font = Font(name="Calibri", size=10, bold=True, color=DETERMINATION_FONTS[val])
        if data:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(data) + 1}"

    # ── Summary ────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = "10B981"
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 46
    due, overdue, closing = deadline_lists(register)
    age = data_age_days(health)
    summary_rows = [
        ("ACCC Acquisitions Register", None, "title"),
        ("", None, None),
        ("Last successful update", fmt_canberra(health.get("last_success_utc") or run_utc), "label"),
        ("Data age", "Up to date" if not age else f"{age} day(s) old", "label"),
        ("Run schedule", SCHEDULE_TEXT, "label"),
        ("", None, None),
        ("REGISTER", None, "section"),
        ("Total matters", len(register), "value"),
        ("Notifications", sum(1 for e in register.values() if e.get("type") == "Notification"), "value"),
        ("Waivers", sum(1 for e in register.values() if e.get("type") == "Waiver"), "value"),
        ("Under assessment", sum(1 for e in register.values() if e.get("acquisition_status") == "Under assessment"), "value"),
        ("", None, None),
        ("DEADLINES", None, "section"),
        ("Due within 10 business days", len(due), "value"),
        ("Past deadline, awaiting publication", len(overdue), "value"),
        ("Consultations closing within 5 business days", len(closing), "value"),
        ("", None, None),
        ("THIS RUN", None, "section"),
        ("New matters", counts["new"], "value"),
        ("Matters with material changes", counts["changed"], "value"),
        ("Removed matters", counts["removed"], "value"),
        ("New decision documents", counts["decisions"], "value"),
        ("Consultation documents downloaded", counts["documents"], "value"),
    ]
    for ri, (label, value, kind) in enumerate(summary_rows, 1):
        a = ws.cell(row=ri, column=1, value=label or None)
        b = ws.cell(row=ri, column=2, value=value)
        ws.row_dimensions[ri].height = 20
        if kind == "title":
            a.font = Font(bold=True, name="Calibri", size=15, color=NAVY)
            ws.row_dimensions[ri].height = 26
        elif kind == "section":
            for c in (a, b):
                c.fill = PatternFill("solid", fgColor=NAVY)
                c.font = Font(bold=True, name="Calibri", size=10, color="FFFFFF")
        else:
            a.font = Font(bold=kind == "label", name="Calibri", size=10)
            b.font = Font(name="Calibri", size=10, bold=isinstance(value, int) and value > 0,
                          color="065F46" if isinstance(value, int) and value > 0 else "111827")
            b.alignment = Alignment(horizontal="left", wrap_text=True)

    # ── Active Matters (deadline tracker) ──────────────────────────────────
    active = [r for r in rows if r["Acquisition Status"] == "Under assessment"]
    order = {"Past deadline (awaiting publication)": 0, "Due today": 1, "Due within 5 business days": 2,
             "Due within 10 business days": 3, "On track": 4}
    active.sort(key=lambda r: (order.get(r["Deadline Status"], 9),
                               r["Business Days Remaining"] if r["Business Days Remaining"] != "" else 999))
    for r in active:
        r["Period"] = f"Day {r['Business Days Elapsed']} of {r['Determination Period Business Days (ACT)']}" \
            if r["Business Days Elapsed"] != "" else ""
    write_table(
        wb.create_sheet("Active Matters"),
        ["Acquisition Name", "Case Number", "Type", "Stage", "Notification / Application Date",
         "End of Determination Period", "Period", "Business Days Remaining", "Deadline Status",
         "Consultation Closes", "Latest Key Event", "Acquirer(s)", "Target(s) / Vendor(s)",
         "M&A Sector", "Watchlist Match", "URL"],
        active, "DC2626",
        widths={"Acquisition Name": 44, "Stage": 24, "Deadline Status": 30, "Latest Key Event": 34,
                "Acquirer(s)": 40, "Target(s) / Vendor(s)": 40, "M&A Sector": 30, "Watchlist Match": 24,
                "Notification / Application Date": 14, "End of Determination Period": 14, "Period": 14},
        wrap_cols={"Acquisition Name", "Latest Key Event", "Acquirer(s)", "Target(s) / Vendor(s)",
                   "M&A Sector", "Stage"},
        date_cols={"Notification / Application Date", "End of Determination Period", "Consultation Closes"},
        link_cols={"URL"}, status_col="Deadline Status")

    # ── Notifications / Waivers ────────────────────────────────────────────
    list_cols = ["Acquisition Name", "Case Number", "Stage", "Acquisition Status", "ACCC Determination",
                 "Notification / Application Date", "End of Determination Period",
                 "Determination Publication Date", "Elapsed Business Days to Decision (ACT)",
                 "Acquirer(s)", "Target(s) / Vendor(s)", "M&A Sector", "Latest Key Event", "URL"]
    list_widths = {"Acquisition Name": 44, "Stage": 24, "Acquisition Status": 20, "ACCC Determination": 22,
                   "Acquirer(s)": 40, "Target(s) / Vendor(s)": 40, "M&A Sector": 30, "Latest Key Event": 34,
                   "Notification / Application Date": 14, "End of Determination Period": 14,
                   "Determination Publication Date": 14, "Elapsed Business Days to Decision (ACT)": 14}
    list_wrap = {"Acquisition Name", "Acquirer(s)", "Target(s) / Vendor(s)", "M&A Sector", "Latest Key Event"}
    list_dates = {"Notification / Application Date", "End of Determination Period", "Determination Publication Date"}
    for name, kind, color in (("Notifications", "Notification", BLUE), ("Waivers", "Waiver", "7C3AED")):
        write_table(wb.create_sheet(name), list_cols, [r for r in rows if r["Type"] == kind], color,
                    list_widths, list_wrap, list_dates, {"URL"}, det_col="ACCC Determination")

    # ── Recent Changes / Case History ──────────────────────────────────────
    hist_cols = ["Detected (Canberra time)", "Case Number", "Matter", "Event", "Detail", "URL"]
    hist_widths = {"Detected (Canberra time)": 22, "Matter": 44, "Event": 26, "Detail": 80}
    write_table(wb.create_sheet("Recent Changes"), hist_cols,
                case_history_rows(changelog, register, RECENT_CHANGES_DAYS), "F59E0B",
                hist_widths, {"Matter", "Detail"}, (), {"URL"})
    history = case_history_rows(changelog, register)
    history.sort(key=lambda r: (r["Case Number"], r["when"] or datetime.datetime.min.replace(tzinfo=datetime.timezone.utc)))
    write_table(wb.create_sheet("Case History"), hist_cols, history, "6B7280",
                hist_widths, {"Matter", "Detail"}, (), {"URL"})

    # ── Analytics ──────────────────────────────────────────────────────────
    ws = wb.create_sheet("Analytics")
    ws.sheet_properties.tabColor = "0EA5E9"
    r = 1

    def block(title, columns, data):
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12, color=NAVY)
        r += 1
        for ci, col in enumerate(columns, 1):
            c = ws.cell(row=r, column=ci, value=col)
            c.fill, c.font = header_fill, header_font
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        r += 1
        for rec in data:
            for ci, val in enumerate(rec, 1):
                c = ws.cell(row=r, column=ci, value=val)
                c.font, c.border = body_font, border
            r += 1
        r += 1

    stat_cols = ["Determined", "Approved", "With conditions", "Not approved", "Approval rate %",
                 "Median bus. days", "Mean bus. days"]

    def stat_vals(s):
        return [s["determined"], s["approved"], s["approved_with_conditions"], s["not_approved"],
                s["approval_rate_pct"], s["median_business_days"], s["mean_business_days"]]

    block("By year", ["Year", "Type", "Filed", "Completed"] + stat_cols,
          [[y["year"], y["type"], y["filed"], y["completed"]] + stat_vals(y) for y in analytics["yearly"]])
    block("By notification stage", ["Stage", "Matters"] + stat_cols,
          [[s, v["matters"]] + stat_vals(v) for s, v in analytics["by_stage"].items()])
    block("By M&A sector", ["Sector", "Matters", "Under assessment"] + stat_cols,
          [[s["sector"], s["matters"], s["under_assessment"]] + stat_vals(s) for s in analytics["by_sector"]])
    block("By month", ["Year", "Month", "Type", "Filed", "Completed"] + stat_cols,
          [[m["year"], datetime.date(2000, m["month"], 1).strftime("%B"), m["type"], m["filed"],
            m["completed"]] + stat_vals(m) for m in analytics["monthly"]])
    ws.column_dimensions["A"].width = 36
    for col in "BCDEFGHIJKL":
        ws.column_dimensions[col].width = 14

    # ── Raw Data ───────────────────────────────────────────────────────────
    raw_cols = list(rows[0].keys()) if rows else []
    write_table(wb.create_sheet("Raw Data"), raw_cols, rows, NAVY,
                {c: 18 for c in raw_cols} | {"Acquisition Name": 44, "Acquirer(s)": 40,
                                             "Target(s) / Vendor(s)": 40, "Other Parties": 40},
                (), {"Notification / Application Date", "End of Determination Period",
                     "Determination Publication Date", "Consultation Closes"},
                {"URL"}, status_col="Deadline Status", det_col="ACCC Determination")

    wb.save(REGISTER_XLSX)
    print(f"  Excel exported: {len(register)} rows → {REGISTER_XLSX}")


# ---------------------------------------------------------------------------
# status.csv (Google Sheets "Last Updated" tab)
# ---------------------------------------------------------------------------

def write_status_csv(run_utc, register, stored_register, new_slugs, changed_slugs,
                     removed_slugs, changelog, this_run_events, health, decision_alerts):
    """Write status.csv for the Last Updated tab in Google Sheets."""
    blank = ["", "", ""]
    age = data_age_days(health)
    rows = [
        ["Metric", "Value", "Notes"],
        blank,
        ["Last successful update", fmt_canberra(health.get("last_success_utc") or run_utc),
         "Up to date" if not age else f"WARNING: data is {age} day(s) old"],
        ["Last run attempt", fmt_canberra(health.get("last_attempt_utc") or run_utc), health.get("last_status", "")],
        ["HTTP requests used", str(health.get("last_request_count", "")), ""],
        blank,
        ["Total entries monitored", str(len(register)), ""],
        ["New entries this run", str(len(new_slugs)), ""],
        ["Changed entries this run", str(len(changed_slugs)), "material changes only"],
        ["Removed entries this run", str(len(removed_slugs)), ""],
        ["New decision documents", str(len(decision_alerts)), ""],
        blank,
        ["Auto-run schedule", SCHEDULE_TEXT, ""],
        blank,
        ["CHANGES IN THIS RUN", "", ""],
        blank,
    ]

    this_run = [c for c in this_run_events if c.get("material", True)]
    if not this_run:
        rows += [["No changes detected in this run.", "", ""], blank]
    for label, kind, sign in (("NEW ENTRIES", "NEW_ENTRY", "+ "), ("CHANGED ENTRIES", "CHANGED", "~ "),
                              ("REMOVED ENTRIES", "REMOVED", "- ")):
        events = [c for c in this_run if c["event"] == kind]
        if not events:
            continue
        rows.append([f"{label} ({len(events)})", "Case Number" if kind != "CHANGED" else "Change", "URL"])
        for c in events:
            entry = register.get(c["slug"]) or stored_register.get(c["slug"], {})
            rows.append([sign + c.get("title", c["slug"]), entry.get("case_number", ""), entry.get("url", "")])
            for ch in c.get("changes", []):
                if ch.get("material", True):
                    rows.append(["", describe_change(ch), ""])
        rows.append(blank)

    rows += [blank, ["CHANGES IN THE LAST 7 DAYS", "", ""], blank]
    week = case_history_rows(changelog, register, 7)
    if not week:
        rows += [["No changes recorded in the last 7 days.", "", ""], blank]
    else:
        rows.append(["Matter", "Change", "Detected"])
        for r in week:
            rows.append([f"{r['Matter']} ({r['Case Number']})", r["Detail"], r["Detected (Canberra time)"]])

    with open(STATUS_CSV, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(rows)
    print(f"  Status CSV written: {STATUS_CSV}")


# ---------------------------------------------------------------------------
# HTML emails
# ---------------------------------------------------------------------------
# Inline styles only: email clients ignore <style> blocks inconsistently.
C_NAVY, C_BLUE, C_TEXT, C_MUTED, C_LINE = "#0D1B2A", "#1D4ED8", "#111827", "#4B5563", "#E5E7EB"
TONES = {
    "red":   ("#FEE2E2", "#991B1B"),
    "amber": ("#FEF3C7", "#92400E"),
    "green": ("#ECFDF5", "#065F46"),
    "blue":  ("#EFF6FF", "#1E40AF"),
    "grey":  ("#F3F4F6", "#374151"),
}


def chip(text, tone="grey"):
    bg, fg = TONES[tone]
    return (f'<span style="display:inline-block;padding:2px 8px;border-radius:10px;'
            f'background:{bg};color:{fg};font-size:12px;font-weight:600;white-space:nowrap">{esc(text)}</span>')


def link(text, url):
    if not url:
        return esc(text)
    return f'<a href="{esc(url)}" style="color:{C_BLUE};text-decoration:none">{esc(text)}</a>'


def stat_tiles(items):
    cells = "".join(
        f'<td style="padding:10px 14px;border:1px solid {C_LINE};border-radius:6px;text-align:center">'
        f'<div style="font-size:22px;font-weight:700;color:{C_NAVY}">{esc(n)}</div>'
        f'<div style="font-size:12px;color:{C_MUTED}">{esc(label)}</div></td>'
        f'<td style="width:8px"></td>'
        for label, n in items)
    return f'<table role="presentation" style="border-collapse:separate;margin:4px 0 8px"><tr>{cells}</tr></table>'


def table(headers, rows):
    head = "".join(f'<th style="text-align:left;padding:7px 8px;font-size:12px;color:{C_MUTED};'
                   f'border-bottom:2px solid {C_LINE};font-weight:600">{esc(h)}</th>' for h in headers)
    body = ""
    for i, row in enumerate(rows):
        bg = "#F9FAFB" if i % 2 else "#FFFFFF"
        body += "<tr>" + "".join(
            f'<td style="padding:7px 8px;font-size:13px;color:{C_TEXT};border-bottom:1px solid {C_LINE};'
            f'vertical-align:top;background:{bg}">{cell}</td>' for cell in row) + "</tr>"
    return (f'<table role="presentation" style="width:100%;border-collapse:collapse;margin:6px 0 4px">'
            f"<tr>{head}</tr>{body}</table>")


def section(title, inner, count=None):
    badge = f' <span style="color:{C_MUTED};font-weight:400">({count})</span>' if count is not None else ""
    return (f'<h2 style="font-size:16px;color:{C_NAVY};margin:26px 0 6px;padding-bottom:4px;'
            f'border-bottom:1px solid {C_LINE}">{esc(title)}{badge}</h2>{inner}')


def banner(text, tone="amber"):
    bg, fg = TONES[tone]
    return (f'<div style="background:{bg};color:{fg};padding:10px 14px;border-radius:6px;'
            f'font-size:14px;margin:12px 0">{text}</div>')


def html_page(title, subtitle, body, footer_links=()):
    links = " &nbsp;·&nbsp; ".join(link(t, u) for t, u in footer_links if u)
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{esc(title)}</title></head>
<body style="margin:0;padding:0;background:#F3F4F6;font-family:Segoe UI,Helvetica,Arial,sans-serif;color:{C_TEXT}">
<table role="presentation" style="width:100%;border-collapse:collapse"><tr><td align="center" style="padding:16px 8px">
<table role="presentation" style="width:100%;max-width:780px;border-collapse:collapse;background:#FFFFFF;border-radius:8px;overflow:hidden">
<tr><td style="background:{C_NAVY};padding:18px 22px">
  <div style="font-size:12px;letter-spacing:.06em;text-transform:uppercase;color:#93C5FD">ACCC Acquisitions Register Monitor</div>
  <div style="font-size:20px;font-weight:700;color:#FFFFFF;margin-top:4px">{esc(title)}</div>
  <div style="font-size:13px;color:#CBD5E1;margin-top:4px">{esc(subtitle)}</div>
</td></tr>
<tr><td style="padding:8px 22px 22px">{body}</td></tr>
<tr><td style="padding:14px 22px;border-top:1px solid {C_LINE};font-size:12px;color:{C_MUTED}">
  {links}<br><span style="color:#9CA3AF">Day counts treat the notification date as Day 0 and use ACT public holidays
  plus the ACCC 23 Dec to 11 Jan suspension. Schedule: {esc(SCHEDULE_TEXT)}</span>
</td></tr></table></td></tr></table></body></html>"""


def _parties_line(entry):
    from scraper import party_names
    acq = ", ".join(party_names(entry.get("acquirers", ""))[:2])
    tgt = ", ".join(party_names(entry.get("targets", ""))[:2])
    if acq and tgt:
        return f"{esc(short(acq, 70))} → {esc(short(tgt, 70))}"
    return esc(short(acq or tgt, 140))


def _matter_cell(entry):
    case = entry.get("case_number", "")
    return f"{link(entry.get('title', ''), entry.get('url', ''))}<br><span style=\"color:{C_MUTED};font-size:12px\">{esc(case)}</span>"


def _deadline_chip(info):
    tone = {"Past deadline (awaiting publication)": "red", "Due today": "red",
            "Due within 5 business days": "amber",
            "Due within 10 business days": "amber", "On track": "green"}.get(info["status"], "grey")
    return chip(info["status"], tone) if info["status"] else ""


def _footer_links():
    return [
        ("Download CSV", repo_url("/raw/main/data/register.csv")),
        ("Full changelog", repo_url("/blob/main/data/changelog.json")),
        ("Run log", run_url()),
        ("ACCC register", REGISTER_URL),
    ]


def _write_outbox(name, content):
    os.makedirs(OUTBOX_DIR, exist_ok=True)
    path = os.path.join(OUTBOX_DIR, name)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path


def _deadline_sections(register, limit=25):
    due, overdue, closing = deadline_lists(register)
    parts = []
    if overdue or due:
        rows = [[_matter_cell(e), esc(display_date(e.get("end_of_determination_period"))),
                 esc(info["label"]), _deadline_chip(info)] for _, e, info in (overdue + due)[:limit]]
        parts.append(section("Decisions due", table(["Matter", "End of period", "Progress", "Status"], rows),
                             len(overdue) + len(due)))
    if closing:
        rows = [[_matter_cell(e), esc(display_date(close.isoformat())),
                 chip(f"{left} business day{'s' if left != 1 else ''} left", "amber" if left <= 2 else "blue")]
                for _, e, close, left in closing[:limit]]
        parts.append(section("Consultations closing soon", table(["Matter", "Closes", ""], rows), len(closing)))
    return "".join(parts)


def _decision_rows(alerts):
    rows = []
    for a in alerts:
        tone = "red" if a["priority"] else "grey"
        doc = link(a["description"], a["url"]) if a["url"] else esc(a["description"])
        rows.append([esc(display_date(a["date"])),
                     f"{link(a['title'], a['accc_url'])}<br><span style=\"color:{C_MUTED};font-size:12px\">{esc(a['case_number'])}</span>",
                     chip(a["category"], tone), doc])
    return rows


def build_changes_email(run_utc, register, stored_register, this_run_events, decision_alerts,
                        downloaded_docs, watch_hits, stats):
    new_events = [e for e in this_run_events if e["event"] == "NEW_ENTRY"]
    changed = [e for e in this_run_events if e["event"] == "CHANGED" and e.get("material")]
    removed = [e for e in this_run_events if e["event"] == "REMOVED"]
    priority = [a for a in decision_alerts if a["priority"]]
    other_decisions = [a for a in decision_alerts if not a["priority"]]

    body = stat_tiles([("New matters", len(new_events)), ("Changed", len(changed)),
                       ("Removed", len(removed)), ("New decisions/events", len(decision_alerts)),
                       ("Documents saved", len(downloaded_docs))])
    if stats.get("schema_upgrade"):
        body += banner("This run upgraded the stored data format (party identifiers, consultation text). "
                       "Formatting-only differences were recorded as cosmetic and are not listed here.", "blue")
    if watch_hits:
        rows = [[_matter_cell(register.get(s, {})), esc("; ".join(r))] for s, r in watch_hits.items()]
        body += section("Watchlist matches", table(["Matter", "Matched on"], rows), len(watch_hits))
    if priority:
        body += section("Priority decisions and events",
                        table(["Date", "Matter", "Event", "Document"], _decision_rows(priority)), len(priority))
    if new_events:
        rows = []
        for ev in new_events:
            e = register.get(ev["slug"], {})
            sectors = ", ".join(get_ma_sectors_for_entry(e.get("anzsic_codes", ""))[:2])
            rows.append([_matter_cell(e), esc(e.get("type", "")), esc(display_date(e.get("notification_date"))),
                         _parties_line(e), esc(sectors)])
        body += section("New matters", table(["Matter", "Type", "Notified", "Parties", "Sector"], rows), len(rows))
    if changed:
        rows = []
        for ev in changed:
            e = register.get(ev["slug"], {})
            items = "".join(f"<li style=\"margin:2px 0\">{esc(describe_change(ch))}</li>"
                            for ch in ev["changes"] if ch.get("material"))
            rows.append([_matter_cell(e), f'<ul style="margin:0;padding-left:18px">{items}</ul>'])
        body += section("Changes", table(["Matter", "What changed"], rows), len(rows))
    if other_decisions:
        body += section("Other decisions and key events",
                        table(["Date", "Matter", "Event", "Document"], _decision_rows(other_decisions)),
                        len(other_decisions))
    if removed:
        rows = [[esc(ev.get("title", "")), esc(ev.get("case_number", ""))] for ev in removed]
        body += section("Removed from the register", table(["Matter", "Case"], rows), len(rows))
    body += _deadline_sections(register)
    body += (f'<p style="font-size:12px;color:{C_MUTED};margin-top:22px">The full register is attached as an Excel '
             f"workbook (Summary, Active Matters, Notifications, Waivers, Recent Changes, Case History, "
             f"Analytics, Raw Data). This run used {stats.get('requests', 0)} requests and re-checked "
             f"{stats.get('fetched', 0)} of {stats.get('listing', 0)} case pages.</p>")

    parts = [f"{len(new_events)} new", f"{len(changed)} changed"]
    if removed:
        parts.append(f"{len(removed)} removed")
    if decision_alerts:
        parts.append(f"{len(decision_alerts)} decision/event")
    subject = "ACCC Register: " + ", ".join(parts)
    if priority:
        subject = "[Priority] " + subject
    _write_outbox("email_changes.html",
                  html_page("Register update", f"Run completed {fmt_canberra(run_utc)}", body, _footer_links()))
    return subject


def build_documents_email(run_utc, downloaded_docs, register):
    attached, linked = documents.stage_attachments(downloaded_docs)
    fresh = [d for d in downloaded_docs if d.get("source") != "changelog backfill"]
    recovered = [d for d in downloaded_docs if d.get("source") == "changelog backfill"]
    attached_names = {d["filename"] for d in attached}

    def rows_for(docs):
        rows = []
        for d in docs:
            e = register.get(d.get("slug"), {})
            close, left = consultation_info(e)
            closes = esc(display_date(close.isoformat())) if close else ""
            if left is not None:
                closes += "<br>" + chip(f"{left} business days left", "amber" if left <= 2 else "blue")
            where = "attached" if d["filename"] in attached_names else link("in repository",
                                                                            repo_url("/tree/main/data/consultation_docs"))
            rows.append([f"{link(d.get('title', ''), d.get('accc_url', ''))}<br><span style=\"color:{C_MUTED};font-size:12px\">{esc(d.get('case_number', ''))}</span>",
                         chip(d.get("kind", "Document"), "blue"),
                         f"{esc(d.get('description') or d['filename'])}<br><span style=\"color:{C_MUTED};font-size:12px\">{esc(d['filename'])} · {esc(d.get('size_kb', ''))} KB · {where}</span>",
                         closes])
        return rows

    body = ""
    if fresh:
        body += section("New consultation documents",
                        table(["Matter", "Kind", "Document", "Consultation closes"], rows_for(fresh)), len(fresh))
    if recovered:
        body += banner("The documents below were published earlier but never downloaded by the previous "
                       "version of this tool (it only collected .docx questionnaires). They have now been recovered.",
                       "blue")
        body += section("Recovered documents",
                        table(["Matter", "Kind", "Document", "Consultation closes"], rows_for(recovered)),
                        len(recovered))
    if linked:
        body += (f'<p style="font-size:12px;color:{C_MUTED}">{len(linked)} file(s) were too large to attach '
                 f"within the email size limit and are linked instead.</p>")
    subject = f"ACCC: {len(downloaded_docs)} consultation document(s)"
    if fresh:
        subject = f"ACCC: {len(fresh)} new consultation document(s) published"
    _write_outbox("email_documents.html",
                  html_page("Consultation documents", f"Run completed {fmt_canberra(run_utc)}", body,
                            [("All stored documents", repo_url("/tree/main/data/consultation_docs")),
                             ("ACCC register", REGISTER_URL)]))
    return subject, len(attached)


def build_watchlist_email(run_utc, register, watch_hits, this_run_events, decision_alerts):
    events_by_slug = {}
    for ev in this_run_events:
        events_by_slug.setdefault(ev["slug"], []).append(ev)
    rows = []
    for slug, reasons in watch_hits.items():
        e = register.get(slug, {})
        lines = []
        for ev in events_by_slug.get(slug, []):
            if ev["event"] == "NEW_ENTRY":
                lines.append("New matter on the register")
            for ch in ev.get("changes", []):
                if ch.get("material"):
                    lines.append(describe_change(ch))
        for a in decision_alerts:
            if a["slug"] == slug:
                lines.append(f"{a['category']}: {a['description']}")
        items = "".join(f"<li>{esc(l)}</li>" for l in dict.fromkeys(lines))
        rows.append([_matter_cell(e), esc("; ".join(reasons)), f'<ul style="margin:0;padding-left:18px">{items}</ul>',
                     _deadline_chip(deadline_info(e))])
    body = section("Watched matters with activity", table(["Matter", "Matched on", "Activity", "Deadline"], rows),
                   len(rows))
    body += (f'<p style="font-size:12px;color:{C_MUTED}">Edit <code>watchlist.json</code> in the repository to '
             f"change the companies, sectors, ANZSIC codes and keywords you watch.</p>")
    subject = f"ACCC watchlist alert: {len(rows)} matter(s)"
    _write_outbox("email_watchlist.html",
                  html_page("Watchlist alert", f"Run completed {fmt_canberra(run_utc)}", body,
                            [("Edit watchlist", repo_url("/blob/main/watchlist.json")), ("ACCC register", REGISTER_URL)]))
    return subject


def digest_due(health, today=None):
    """The weekly digest goes out on the first successful run of each ISO week."""
    today = today or canberra_today()
    year, week, _ = today.isocalendar()
    key = f"{year}-W{week:02d}"
    return health.get("last_digest_week") != key, key


def build_digest_email(run_utc, register, changelog, analytics):
    week_ago = utc_now() - datetime.timedelta(days=7)
    recent = [ev for ev in changelog if (parse_iso(ev.get("timestamp_utc")) or week_ago) > week_ago]
    new_slugs = [ev["slug"] for ev in recent if ev["event"] == "NEW_ENTRY"]
    decided = [e for e in register.values()
               if (parse_date(e.get("determination_publication_date")) or datetime.date.min)
               >= (canberra_today() - datetime.timedelta(days=7))]
    events = []
    for e in register.values():
        for d in e.get("decisions_docs", []) or []:
            when = parse_date(d.get("date"))
            if when and when >= canberra_today() - datetime.timedelta(days=7):
                category, priority = documents.classify_decision_document(d.get("description", ""))
                if category not in ("Waiver determination", "Phase 1 determination") or priority:
                    events.append({**d, "title": e.get("title"), "case_number": e.get("case_number"),
                                   "accc_url": e.get("url"), "category": category, "priority": priority})
    events.sort(key=lambda a: parse_date(a["date"]) or datetime.date.min, reverse=True)

    new_n = [s for s in new_slugs if register.get(s, {}).get("type") == "Notification"]
    new_w = [s for s in new_slugs if register.get(s, {}).get("type") == "Waiver"]
    approved = sum(1 for e in decided if e.get("accc_determination") in APPROVED)
    not_approved = sum(1 for e in decided if e.get("accc_determination") == "Not approved")
    active = sum(1 for e in register.values() if e.get("acquisition_status") == "Under assessment")
    body = stat_tiles([("New notifications", len(new_n)), ("New waivers", len(new_w)),
                       ("Decided", len(decided)), ("Not approved", not_approved), ("Under assessment", active)])

    if new_n:
        rows = [[_matter_cell(register[s]), esc(display_date(register[s].get("notification_date"))),
                 _parties_line(register[s]),
                 esc(", ".join(get_ma_sectors_for_entry(register[s].get("anzsic_codes", ""))[:2]))]
                for s in new_n]
        body += section("New notifications this week", table(["Matter", "Notified", "Parties", "Sector"], rows), len(rows))
    if decided:
        rows = []
        for e in sorted(decided, key=lambda e: parse_date(e.get("determination_publication_date")), reverse=True):
            det = e.get("accc_determination", "")
            tone = "red" if det == "Not approved" else "amber" if det == "Approved with conditions" else "green"
            rows.append([_matter_cell(e), esc(e.get("type", "")),
                         chip(det or e.get("acquisition_status", ""), tone),
                         esc(e.get("elapsed_business_days", "")),
                         esc(display_date(e.get("determination_publication_date")))])
        body += section("Determinations published this week",
                        table(["Matter", "Type", "Outcome", "Business days", "Published"], rows[:60]), len(rows))
    if events:
        rows = [[esc(display_date(a["date"])), f"{link(a['title'], a['accc_url'])}",
                 chip(a["category"], "red" if a["priority"] else "grey"),
                 link(a["description"], a.get("url", ""))] for a in events[:40]]
        body += section("Notable key events this week", table(["Date", "Matter", "Event", "Document"], rows), len(events))
    body += _deadline_sections(register)

    n = analytics["by_type"].get("Notification", {})
    w = analytics["by_type"].get("Waiver", {})
    body += section("Register at a glance", table(
        ["", "Matters", "Approval rate", "Median business days to decision"],
        [["Notifications", esc(n.get("matters", 0)), esc(f"{n.get('approval_rate_pct')}%" if n.get("approval_rate_pct") is not None else "-"),
          esc(n.get("median_business_days") or "-")],
         ["Waivers", esc(w.get("matters", 0)), esc(f"{w.get('approval_rate_pct')}%" if w.get("approval_rate_pct") is not None else "-"),
          esc(w.get("median_business_days") or "-")]]))

    subject = (f"ACCC weekly digest: {len(new_n)} new notification(s), {len(decided)} decided, "
               f"{approved} approved")
    _write_outbox("email_digest.html",
                  html_page("Weekly digest", f"Week to {display_date(canberra_today().isoformat())}", body,
                            _footer_links()))
    return subject


def failure_subject(health, kind):
    n = health.get("consecutive_failures", 1)
    what = {"blocked": "ACCC website is blocking the monitor",
            "sanity": "data check failed, nothing was saved",
            "error": "the scraper crashed"}.get(kind, "run failed")
    return f"ACCC Monitor FAILED ({n} in a row): {what}"


def write_failure_email(health, kind, message):
    age = data_age_days(health)
    last = health.get("last_success_utc", "")
    advice = {
        "blocked": (
            "<p>The ACCC website's CDN (Akamai) answered <b>403 Access Denied</b>. The refusal happens before the "
            "page is served and also applies to <code>robots.txt</code>, so it is based on where the request comes "
            "from (cloud and datacenter IP ranges), not on the pages requested.</p>"
            "<p>Options, in order of effort:</p><ol>"
            "<li>Run the monitor from a machine on a normal business or home connection by registering it as a "
            "GitHub <b>self-hosted runner</b>, then set the repository variable <code>RUNNER_LABEL</code> to "
            "<code>self-hosted</code>. No code change needed.</li>"
            "<li>Add a contact to the User-Agent with the repository variable <code>MONITOR_CONTACT</code> "
            "(an email address or URL) and ask the ACCC to allow the monitor.</li></ol>"
            "<p>The monitor already keeps its request volume low and stops after a few refusals, so it will not "
            "keep hitting the site while blocked.</p>"),
        "sanity": (
            "<p>The scraper refused to save because the data it received looked wrong (too few register entries, "
            "or key fields missing on many pages). This usually means the ACCC changed its page layout. Stored "
            "data was left untouched. The parser in <code>scraper.py</code> (parse_detail_page / parse_card) "
            "needs checking against a live page.</p>"),
        "error": "<p>The scraper stopped with an unexpected error. The run log has the full traceback.</p>",
    }.get(kind, "")
    body = banner(f"<b>{esc(message)}</b>", "red")
    body += table(["", ""], [
        ["Consecutive failed runs", esc(health.get("consecutive_failures", 1))],
        ["Last successful update", esc(fmt_canberra(last) or "never")],
        ["Data age", esc(f"{age} day(s)" if age is not None else "unknown")],
    ])
    body += section("What this means", advice)
    body += ('<p style="font-size:13px">Until this is fixed, register data, deadlines and alerts are not '
             "being updated. You will be reminded about once a day while it keeps failing.</p>")
    _write_outbox("email_failure.html",
                  html_page("Monitor run failed", f"Attempted {fmt_canberra(health.get('last_attempt_utc'))}", body,
                            [("Run log", run_url()), ("Repository", repo_url())]))


# ---------------------------------------------------------------------------
# latest_changes.json
# ---------------------------------------------------------------------------

def write_latest_changes(run_utc, register, stored_register, this_run_events, decision_alerts):
    def entry_ref(slug, source):
        e = source.get(slug, {})
        return {"slug": slug, "title": e.get("title", slug), "case_number": e.get("case_number", ""),
                "url": e.get("url", "")}

    changed = [ev for ev in this_run_events if ev["event"] == "CHANGED"]
    save_json(LATEST_CHANGES_JSON, {
        "run_utc": run_utc,
        "run_canberra": fmt_canberra(run_utc),
        "new_count": sum(1 for ev in this_run_events if ev["event"] == "NEW_ENTRY"),
        "changed_count": sum(1 for ev in changed if ev.get("material")),
        "cosmetic_only_count": sum(1 for ev in changed if not ev.get("material")),
        "removed_count": sum(1 for ev in this_run_events if ev["event"] == "REMOVED"),
        "new_entries": [entry_ref(ev["slug"], register) for ev in this_run_events if ev["event"] == "NEW_ENTRY"],
        "changed_entries": [
            {**entry_ref(ev["slug"], register),
             "material": ev.get("material", True),
             "summary": [describe_change(ch) for ch in ev["changes"] if ch.get("material", True)],
             "changes": ev["changes"]}
            for ev in changed
        ],
        "removed_entries": [entry_ref(ev["slug"], stored_register) for ev in this_run_events if ev["event"] == "REMOVED"],
        "decision_documents": decision_alerts,
    })


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def holiday_coverage_warning(register):
    latest = max((d for e in register.values() for f in DATE_FIELDS
                  for d in [parse_date(e.get(f, ""))] if d), default=None)
    if latest and latest.year > HOLIDAY_TABLE_LAST_YEAR:
        return (f"ACT holiday table ends in {HOLIDAY_TABLE_LAST_YEAR} but register dates reach {latest}. "
                "Business day counts after that year ignore public holidays.")
    return ""


def build_all_outputs(run_utc, register, stored_register, changelog, this_run_events, new_slugs,
                      changed_slugs, removed_slugs, downloaded_docs, manifest, decision_alerts, health, stats):
    watchlist = load_watchlist()

    rows = register_rows(register, manifest, watchlist)
    export_csv(rows)

    summary = generate_summary(register, health, manifest)
    warning = holiday_coverage_warning(register)
    if warning:
        summary["warnings"] = [warning]
        print(f"  WARNING: {warning}")
    save_json(SUMMARY_JSON, summary)

    write_latest_changes(run_utc, register, stored_register, this_run_events, decision_alerts)
    write_status_csv(run_utc, register, stored_register, new_slugs, changed_slugs, removed_slugs,
                     changelog, this_run_events, health, decision_alerts)
    export_excel(register, rows, changelog, summary["analytics"], health, run_utc, {
        "new": len(new_slugs), "changed": len(changed_slugs), "removed": len(removed_slugs),
        "decisions": len(decision_alerts), "documents": len(downloaded_docs),
    })

    # Watchlist hits among this run's activity
    active_slugs = set(new_slugs) | set(changed_slugs) | {a["slug"] for a in decision_alerts}
    watch_hits = {}
    for slug in active_slugs:
        reasons = watchlist_matches(register.get(slug, {}), watchlist)
        if reasons:
            watch_hits[slug] = reasons

    changes_subject = build_changes_email(run_utc, register, stored_register, this_run_events,
                                          decision_alerts, downloaded_docs, watch_hits, stats)
    documents_subject, attached = ("", 0)
    if downloaded_docs:
        documents_subject, attached = build_documents_email(run_utc, downloaded_docs, register)
    watchlist_subject = ""
    if watch_hits:
        watchlist_subject = build_watchlist_email(run_utc, register, watch_hits, this_run_events, decision_alerts)

    send_digest, week_key = digest_due(health)
    digest_subject = ""
    if send_digest:
        digest_subject = build_digest_email(run_utc, register, changelog, summary["analytics"])
        health["last_digest_week"] = week_key

    return {
        "run_local": fmt_canberra(run_utc),
        "changes_subject": changes_subject,
        "documents_email": bool(downloaded_docs),
        "documents_subject": documents_subject,
        "attached_doc_count": attached,
        "watchlist_email": bool(watch_hits),
        "watchlist_subject": watchlist_subject,
        "digest_email": send_digest,
        "digest_subject": digest_subject,
    }
